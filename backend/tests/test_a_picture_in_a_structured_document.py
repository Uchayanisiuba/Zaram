"""A picture the person attached can go on an invoice, a CV and a slide deck.

`test_an_attached_picture_becomes_content.py` covers a document written as
markdown, where the model puts `![alt](photo-1.png)` on a line of its own. An
invoice, a CV and a deck are not written that way: they are filled from fields
read out of the answer (`artifacts/extract.py`), so they had nowhere to put a
picture and "put my logo on this invoice" produced an invoice without one.

The picture is *named* in a field -- an invoice's `logo`, a CV's `photo`, a
slide's `image` -- and the bytes are read afterwards on this machine. A
spreadsheet and a chart stay out: cells and a chart drawn from data have no place
a picture belongs.

These go through `DocumentsRuntime.execute`, the call the conversation makes, and
read the file that comes out; and one goes through the dispatcher, because a note
about a picture that was left out is only a note if it reaches the person.
"""

from __future__ import annotations

import asyncio
import base64
import io
import json
import zipfile

import pytest
from PIL import Image

from artifacts.extract import chosen_picture
from artifacts.records import ArtifactRecords
from artifacts.service import ArtifactService
from artifacts.store import ArtifactStore
from attachments import pictures
from attachments.store import AttachmentStore
from runtimes.documents.runtime import GENERATE, DocumentsRuntime


def _png(colour: str = "red", size: int = 24) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (size, size), colour).save(out, format="PNG")
    return out.getvalue()


def _b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


@pytest.fixture
def store(tmp_path):
    held = AttachmentStore(str(tmp_path / "attachments"))
    pictures.use_store(held)
    yield held
    pictures.use_store(None)


INVOICE_JSON = {
    "bill_to": ["Acme Ltd"],
    "currency": "USD",
    "terms_days": 14,
    "items": [{"description": "Design", "quantity": 2, "unit": "days", "unit_price": 400}],
}
CV_JSON = {
    "name": "Adaeze Okonkwo",
    "headline": "Producer",
    "summary": "Ten years producing documentaries.",
    "experience": [{"role": "Producer", "organisation": "Harbour Lane", "dates": "2020 - 2024"}],
}
DECK_JSON = {
    "slides": [
        {"heading": "The problem", "bullets": ["Slow", "Costly"]},
        {"heading": "The fix", "bullets": ["Faster", "Cheaper"]},
    ]
}


class Asker:
    """A stand-in for the model that reads the answer into fields.

    It records every prompt, so the tests can see what the extraction step was
    and was not told about pictures.
    """

    def __init__(self, payload: dict):
        self.payload = payload
        self.prompts: list[str] = []

    def __call__(self, prompt: str, system: str = "") -> str:
        self.prompts.append(prompt)
        return json.dumps(self.payload)


@pytest.fixture
def runtime(tmp_path):
    service = ArtifactService(
        ArtifactRecords(str(tmp_path / "artifacts.db")), ArtifactStore(tmp_path / "out")
    )
    rt = DocumentsRuntime(service)
    rt.set_pictures(pictures.document_pictures_for_session)
    return rt


def make(runtime, asker, prompt, answer, fmt="html", session="s1"):
    runtime.set_extractor(asker)
    return asyncio.run(
        runtime.execute(
            GENERATE,
            {"prompt": prompt, "answer": answer, "session_id": session, "format": fmt},
        )
    )


def read_html(result) -> str:
    with open(result["artifact"]["path"], encoding="utf-8") as handle:
        return handle.read()


def media(result, folder: str) -> list[str]:
    names = zipfile.ZipFile(result["artifact"]["path"]).namelist()
    return [n for n in names if n.startswith(folder)]


ANSWER = "Invoice for Acme Ltd: two days of design at 400 a day, due in 14 days."


# ---- invoice and quote ---------------------------------------------------------


class TestAnInvoiceCarriesALogo:
    def test_the_attached_logo_is_in_the_html_with_its_real_bytes(self, runtime, store):
        data = _png("green")
        item, _ = store.add("s1", "logo.png", data)
        asker = Asker({**INVOICE_JSON, "logo": item.page_name})
        result = make(runtime, asker, "make me an invoice for Acme with my logo", ANSWER)
        assert result["success"], result
        html = read_html(result)
        assert _b64(data) in html
        assert 'class="logo"' in html
        assert "notes" not in result

    def test_a_quote_takes_one_too(self, runtime, store):
        data = _png("blue")
        item, _ = store.add("s1", "logo.png", data)
        asker = Asker({**INVOICE_JSON, "logo": item.page_name})
        result = make(runtime, asker, "write a quote for Acme, put my logo on the quote", ANSWER)
        assert result["success"], result
        assert _b64(data) in read_html(result)

    def test_the_word_file_carries_it_too(self, runtime, store):
        item, _ = store.add("s1", "logo.png", _png("green"))
        asker = Asker({**INVOICE_JSON, "logo": item.page_name})
        result = make(
            runtime, asker, "make me an invoice for Acme with my logo", ANSWER, fmt="docx"
        )
        assert result["success"], result
        assert media(result, "word/media/image"), "the logo is not in the .docx"

    def test_it_is_recorded_as_a_source_the_person_supplied(self, runtime, store):
        item, _ = store.add("s1", "logo.png", _png())
        asker = Asker({**INVOICE_JSON, "logo": item.page_name})
        result = make(runtime, asker, "make me an invoice for Acme with my logo", ANSWER)
        sources = {(s["kind"], s["title"]) for s in result["artifact"]["sources"]}
        assert ("user_document", "logo.png (picture you attached)") in sources

    def test_the_extractor_is_offered_the_names_only_when_pictures_are_attached(
        self, runtime, store
    ):
        asker = Asker(INVOICE_JSON)
        make(runtime, asker, "make me an invoice for Acme", ANSWER)
        assert "attached these pictures" not in asker.prompts[0]
        assert '"logo"' not in asker.prompts[0]

        store.add("s1", "logo.png", _png())
        asker = Asker(INVOICE_JSON)
        make(runtime, asker, "make me an invoice for Acme", ANSWER)
        assert "logo.png" in asker.prompts[0] and '"logo"' in asker.prompts[0]

    def test_a_name_that_is_not_attached_is_left_out_and_said(self, runtime, store):
        store.add("s1", "logo.png", _png())
        asker = Asker({**INVOICE_JSON, "logo": "company-crest.png"})
        result = make(runtime, asker, "make me an invoice for Acme with my logo", ANSWER)
        assert result["success"], result
        html = read_html(result)
        assert 'class="logo"' not in html and "base64" not in html
        assert any("company-crest.png" in n and "left out" in n for n in result["notes"])
        assert not [s for s in result["artifact"]["sources"] if s["kind"] == "user_document"]

    def test_another_conversations_picture_is_not_used(self, runtime, store):
        item, _ = store.add("s2", "logo.png", _png("green"))
        asker = Asker({**INVOICE_JSON, "logo": item.page_name})
        result = make(runtime, asker, "make me an invoice for Acme with my logo", ANSWER)
        assert result["success"], result
        assert "base64" not in read_html(result)
        assert "logo.png" not in asker.prompts[0]
        # Nothing was offered, so nothing was asked: the field is not even in
        # the prompt, and there is no picture to place.
        assert "notes" not in result

    def test_a_picture_attached_for_another_reason_is_not_made_a_logo(self, runtime, store):
        """A receipt attached so a model could read it is not the logo."""
        item, _ = store.add("s1", "receipt.png", _png("red"))
        asker = Asker({**INVOICE_JSON, "logo": item.page_name})
        result = make(runtime, asker, "make me an invoice for Acme", ANSWER)
        assert result["success"], result
        assert "base64" not in read_html(result)
        assert "notes" not in result


# ---- CV ----------------------------------------------------------------------------


CV_ANSWER = "Adaeze Okonkwo, producer at Harbour Lane 2020 to 2024, ten years in documentary."


class TestACvCarriesAPhoto:
    def test_the_attached_photo_is_in_the_html(self, runtime, store):
        data = _png("yellow", 40)
        item, _ = store.add("s1", "me.png", data)
        asker = Asker({**CV_JSON, "photo": item.page_name})
        result = make(runtime, asker, "write my CV with my photo at the top", CV_ANSWER)
        assert result["success"], result
        html = read_html(result)
        assert _b64(data) in html
        assert 'class="photo"' in html
        assert "Photo of Adaeze Okonkwo" in html

    def test_a_cv_without_one_has_the_markup_it_always_had(self, runtime, store):
        asker = Asker(CV_JSON)
        result = make(runtime, asker, "write my CV", CV_ANSWER)
        html = read_html(result)
        assert 'class="cv-head"' not in html

    def test_the_word_file_carries_it_and_it_is_small(self, runtime, store):
        pytest.importorskip("docx")
        from docx import Document
        from docx.shared import Mm

        item, _ = store.add("s1", "me.png", _png("yellow", 600))
        asker = Asker({**CV_JSON, "photo": item.page_name})
        result = make(
            runtime, asker, "write my CV with my photo at the top", CV_ANSWER, fmt="docx"
        )
        assert result["success"], result
        assert media(result, "word/media/image"), "the photo is not in the .docx"
        shapes = Document(result["artifact"]["path"]).inline_shapes
        assert shapes[0].width <= Mm(32) and shapes[0].height <= Mm(40)

    def test_it_is_recorded_as_a_source_the_person_supplied(self, runtime, store):
        item, _ = store.add("s1", "me.png", _png())
        asker = Asker({**CV_JSON, "photo": item.page_name})
        result = make(runtime, asker, "write my CV with my photo at the top", CV_ANSWER)
        sources = {(s["kind"], s["title"]) for s in result["artifact"]["sources"]}
        assert ("user_document", "me.png (picture you attached)") in sources

    def test_an_unknown_name_is_left_out_and_said(self, runtime, store):
        store.add("s1", "me.png", _png())
        asker = Asker({**CV_JSON, "photo": "headshot-2019.png"})
        result = make(runtime, asker, "write my CV with my photo", CV_ANSWER)
        assert result["success"], result
        assert "base64" not in read_html(result)
        assert any("headshot-2019.png" in n for n in result["notes"])

    def test_another_conversations_picture_is_not_used(self, runtime, store):
        item, _ = store.add("s2", "me.png", _png())
        asker = Asker({**CV_JSON, "photo": item.page_name})
        result = make(runtime, asker, "write my CV with my photo", CV_ANSWER)
        assert "base64" not in read_html(result)
        assert "me.png" not in asker.prompts[0]
        assert not [s for s in result["artifact"]["sources"] if s["kind"] == "user_document"]


# ---- deck ----------------------------------------------------------------------------


DECK_ANSWER = "The problem: slow and costly. The fix: faster and cheaper."


class TestADeckCarriesAPicturePerSlide:
    def test_the_picture_is_under_its_slides_heading_in_the_html(self, runtime, store):
        data = _png("purple", 40)
        item, _ = store.add("s1", "diagram.png", data)
        asker = Asker(
            {
                "slides": [
                    DECK_JSON["slides"][0],
                    {**DECK_JSON["slides"][1], "image": item.page_name},
                ]
            }
        )
        result = make(runtime, asker, "make slides, put my diagram on the fix", DECK_ANSWER)
        assert result["success"], result
        html = read_html(result)
        assert _b64(data) in html
        assert html.index("The fix") < html.index(_b64(data)) < html.index("Faster")
        assert html.count(_b64(data)) == 1

    def test_the_pptx_puts_it_on_that_slide_beside_the_bullets(self, runtime, store):
        pytest.importorskip("pptx")
        from pptx import Presentation

        item, _ = store.add("s1", "diagram.png", _png("purple", 400))

        def build(with_picture: bool):
            slides = [dict(s) for s in DECK_JSON["slides"]]
            if with_picture:
                slides[1]["image"] = item.page_name
            asker = Asker({"slides": slides})
            return make(
                runtime,
                asker,
                "make slides, put my diagram on the fix",
                DECK_ANSWER,
                fmt="pptx",
            )

        result = build(True)
        assert result["success"], result
        assert media(result, "ppt/media/image"), "the picture is not in the .pptx"

        deck = Presentation(result["artifact"]["path"])
        carrying = [
            s
            for s in deck.slides
            if any(shape.shape_type == 13 for shape in s.shapes)
        ]
        assert len(carrying) == 1
        slide = carrying[0]
        assert slide.shapes.title.text == "The fix"
        assert any(
            shape.has_text_frame and "Faster" in shape.text_frame.text
            for shape in slide.shapes
        ), "the picture has a slide of its own instead of sharing the slide's"
        # The bullets were narrowed to make room, not left underneath it.
        picture = next(s for s in slide.shapes if s.shape_type == 13)
        body = slide.placeholders[1]
        assert body.left + body.width <= picture.left

        # And it did not add a slide.
        plain = Presentation(build(False)["artifact"]["path"])
        assert len(deck.slides) == len(plain.slides)

    def test_it_is_recorded_as_a_source_once_however_often_it_is_used(self, runtime, store):
        item, _ = store.add("s1", "diagram.png", _png())
        slides = [{**s, "image": item.page_name} for s in DECK_JSON["slides"]]
        result = make(
            runtime, Asker({"slides": slides}), "make slides with my diagram", DECK_ANSWER
        )
        mine = [s for s in result["artifact"]["sources"] if s["kind"] == "user_document"]
        assert [s["title"] for s in mine] == ["diagram.png (picture you attached)"]

    def test_an_unknown_name_is_left_out_and_said(self, runtime, store):
        store.add("s1", "diagram.png", _png())
        slides = [dict(s) for s in DECK_JSON["slides"]]
        slides[0]["image"] = "chart-final.png"
        result = make(
            runtime, Asker({"slides": slides}), "make slides with my diagram", DECK_ANSWER
        )
        assert result["success"], result
        assert "base64" not in read_html(result)
        assert any("chart-final.png" in n for n in result["notes"])

    def test_another_conversations_picture_is_not_used(self, runtime, store):
        item, _ = store.add("s2", "diagram.png", _png())
        slides = [{**s, "image": item.page_name} for s in DECK_JSON["slides"]]
        asker = Asker({"slides": slides})
        result = make(runtime, asker, "make slides with my diagram", DECK_ANSWER)
        assert "base64" not in read_html(result)
        assert "diagram.png" not in asker.prompts[0]
        assert not [s for s in result["artifact"]["sources"] if s["kind"] == "user_document"]


# ---- what each kind is told ------------------------------------------------------


class TestWhatTheModelIsTold:
    def _refs(self, store, *names):
        return pictures.refs_for([store.add("s1", n, _png())[0] for n in names])

    @pytest.mark.parametrize(
        "slot,words",
        [("logo", "Logo:"), ("photo", "Photo:"), ("slide", "Slide image:")],
    )
    def test_each_kind_is_told_the_names_and_its_own_field(self, store, slot, words):
        told = pictures.document_guidance(
            self._refs(store, "logo.png", "me.png"), seen=True, slot=slot
        )
        assert words in told
        assert "logo.png" in told and "me.png" in told
        assert "markdown image" not in told

    def test_a_document_written_as_markdown_is_told_as_before(self, store):
        told = pictures.document_guidance(self._refs(store, "logo.png"), seen=True)
        assert "markdown image" in told and "Logo:" not in told

    def test_nothing_attached_nothing_said(self):
        for slot in ("", "logo", "photo", "slide"):
            assert pictures.document_guidance([], seen=True, slot=slot) == ""

    def test_the_bytes_are_never_in_the_guidance(self, store):
        item, _ = store.add("s1", "logo.png", _png())
        told = pictures.document_guidance(pictures.refs_for([item]), seen=True, slot="logo")
        assert _b64(_png())[:40] not in told

    def test_the_runtime_names_the_field_for_each_request(self, runtime):
        assert runtime.picture_slot("make me an invoice for Acme") == "logo"
        assert runtime.picture_slot("write a quote for Acme") == "logo"
        assert runtime.picture_slot("write my CV") == "photo"
        assert runtime.picture_slot("make me a slide deck") == "slide"
        assert runtime.picture_slot("write a cover letter") == ""

    def test_a_spreadsheet_and_a_chart_are_offered_no_picture(self, runtime):
        assert not runtime.places_pictures("make me a spreadsheet of expenses")
        assert not runtime.places_pictures("draw a chart of sales")
        assert runtime.places_pictures("make me an invoice for Acme")

    def test_the_writing_step_gets_the_field_for_the_kind_being_made(self, store, runtime):
        """The engine's call, with the probe the bootstrapper hands it."""
        from core.execution_engine import ExecutionEngine

        store.add("s1", "logo.png", _png())
        engine = ExecutionEngine.__new__(ExecutionEngine)
        engine._picture_names = lambda session, request: pictures.refs_for_session(session)
        engine._picture_slot = runtime.picture_slot
        told = engine._document_picture_guidance("s1", "make me an invoice for Acme", True)
        assert "Logo:" in told and "logo.png" in told
        told = engine._document_picture_guidance("s1", "write my CV", True)
        assert "Photo:" in told

    def test_a_bootstrapped_engine_wires_the_slot(self):
        """The caller, seen: `set_picture_names` stores both probes."""
        from core.execution_engine import ExecutionEngine

        engine = ExecutionEngine.__new__(ExecutionEngine)
        engine.set_picture_names(lambda s, r: [], lambda r: "logo")
        assert engine._picture_slot("x") == "logo"


class TestAChosenName:
    def test_only_an_attached_name_survives(self):
        left_out: list[str] = []
        assert (
            chosen_picture("LOGO.png", ["logo.png"], request="x", answer="logo.png", left_out=left_out)
            == "logo.png"
        )
        assert (
            chosen_picture("https://x.test/a.png", ["logo.png"], request="logo", answer="", left_out=left_out)
            == ""
        )
        assert len(left_out) == 1 and "https" not in left_out[0].replace("httpsx", "")

    @pytest.mark.parametrize("raw", [None, "", "null", "None", False, "  "])
    def test_nothing_named_is_nothing_placed_and_nothing_reported(self, raw):
        left_out: list[str] = []
        assert chosen_picture(raw, ["logo.png"], request="logo", answer="", left_out=left_out) == ""
        assert left_out == []


# ---- the person is told ----------------------------------------------------------------


class TestTheNoteReachesThePerson:
    def test_the_dispatcher_turns_a_note_into_a_notice(self):
        from core.contracts import ExecutionStep
        from core.dispatcher import ARTIFACT_MARKER, NOTICE_MARKER, ExecutionDispatcher

        class FakeRuntime:
            async def execute(self, capability_id, input_data):
                return {
                    "success": True,
                    "artifact": {"id": "a1"},
                    "notes": ['"crest.png" is not one of the pictures you attached, so it was left out.'],
                }

        class FakeRouter:
            def resolve(self, capability_id):
                return FakeRuntime()

        tokens = list(
            ExecutionDispatcher(FakeRouter()).execute_step(
                ExecutionStep(capability_id="document.generate", input_data={})
            )
        )
        assert tokens[0].startswith(ARTIFACT_MARKER)
        notice = next(t for t in tokens if t.startswith(NOTICE_MARKER))
        assert "crest.png" in json.loads(notice[len(NOTICE_MARKER):])["content"]
