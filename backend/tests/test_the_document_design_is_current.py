"""The blocks that carry design, all the way from a model's markdown to a file.

**The vocabulary was the ceiling, and that is what this suite guards.** A model
asked for a proposal writes markdown, and markdown has headings, paragraphs,
lists and tables — no way at all to say "this line is the point". So every
sentence arrived with equal weight, the deck exporter turned all of them into
bullets, and the output was flat however good the writing was. The composer was
never the problem; the design had nowhere to come from.

Four blocks close that, and each one has to survive five hand-offs: a fence in
the model's reply, a contract, the HTML that is the source of truth, the reader
that parses it back, and two binary formats. A test that stopped at the HTML
would pass while a `.docx` lost the lot — which is the exact failure
`export/_reader.py` records about pictures, found only because somebody opened
the file.

The other half is **degradation**, and it is the reason these are roles on
ordinary tags rather than tags of their own. An exporter that has never heard
of a statement must write the paragraph. Several tests below assert that
directly, because it is the property that makes adding a fifth block safe.
"""

from __future__ import annotations

import io

import pytest

from artifacts import theme
from artifacts.contracts import Callout, Divider, Metric, Statement
from artifacts.export import _reader
from artifacts.html import render_document
from artifacts.markdown_blocks import blocks_from_markdown

MARKDOWN = """## The problem

Ordinary prose, which must stay ordinary prose.

```statement
Reconciliation takes two days a week. It should take twenty minutes.
```

```callout
Week 3 is the one that matters.
```

```divider
What it costs
```

```metric
18 | days
11,160 | total fee
```

```callout warn
Without the historical invoices, week 2 becomes a guess.
```

```python
print("this is a code sample, not a design block")
```
"""


def _blocks():
    return blocks_from_markdown(MARKDOWN, title="A proposal")


def _html():
    return render_document(title="A proposal", blocks=_blocks())


# --------------------------------------------------------------------------- #
# The model's markdown reaches the block types
# --------------------------------------------------------------------------- #


def test_a_fence_becomes_the_block_it_names():
    kinds = [type(block).__name__ for block in _blocks()]
    assert "Statement" in kinds
    assert "Callout" in kinds
    assert "Divider" in kinds
    assert "Metric" in kinds


def test_a_real_code_block_is_still_a_code_block():
    """The risk this syntax accepts, asserted rather than assumed.

    Reusing the fence info string means a language literally called `statement`
    would be captured. No such language exists — but `python` does, and a
    proposal containing a code sample has to keep it.
    """
    html = _html()
    assert "code sample, not a design block" in html
    assert "<code>" in html


def test_an_empty_fence_produces_nothing():
    """A model that opened a block and said nothing.

    A block with no content renders as a rule across an empty page, which reads
    as a fault in the document rather than in the reply.
    """
    blocks = blocks_from_markdown("```statement\n\n```\n", title="T")
    assert not [b for b in blocks if isinstance(b, Statement)]


def test_a_metric_line_without_a_separator_keeps_its_value():
    """The number was written and meant. A caption is what is missing.

    Dropping the line would delete a figure the author put in; keeping it
    without a label is recoverable by whoever reads the document.
    """
    blocks = blocks_from_markdown("```metric\n42\n```\n", title="T")
    metric = next(b for b in blocks if isinstance(b, Metric))
    assert len(metric.items) == 1


def test_an_unknown_tone_falls_back_to_the_calm_one():
    """`tone="critical"` is a reasonable guess at a vocabulary nobody published.

    The answer is the quieter of the two panels — never a traceback, and never
    an invented third appearance that would make the two-tone rule a lie.
    """
    html = render_document(title="T", blocks=[Callout(text="Mind this", tone="critical")])
    assert 'class="callout note"' in html


# --------------------------------------------------------------------------- #
# The HTML carries the design, and the reader carries it back
# --------------------------------------------------------------------------- #


def test_the_html_marks_each_block_with_a_role():
    html = _html()
    assert 'class="statement"' in html
    assert 'class="callout note"' in html
    assert 'class="callout warn"' in html
    assert 'class="divider"' in html
    assert 'class="metrics"' in html


def test_a_statement_is_a_paragraph_and_a_divider_is_a_heading():
    """The degradation guarantee, at the markup level.

    A new tag in the block stream makes every existing consumer render
    something unintended — `Table.after_block`'s note records what that costs.
    A role on a tag everybody already handles costs nothing.
    """
    doc = _reader.read(_html())
    by_role = {block.role: block for block in doc.body_blocks() if block.role}
    assert by_role["statement"].tag == "p"
    assert by_role["callout"].tag == "p"
    assert by_role["divider"].tag == "h2"
    assert by_role["metric"].tag == "p"


def test_the_reader_only_believes_roles_it_knows():
    """An arbitrary class is not an instruction to the exporters.

    Treating any class as a role would turn every future stylesheet hook into
    behaviour, which is `markdown_blocks`'s rule — *nothing may invent markup
    the readers would have to learn* — broken from the other end.
    """
    doc = _reader.read('<html><body><p class="sidebar callout">Hi</p></body></html>')
    block = doc.body_blocks()[0]
    assert block.role == "callout"


def test_a_metric_keeps_its_value_and_its_label_apart():
    """Read off `Run.bold`, which the HTML sets deliberately.

    Splitting the rendered text on a space is the guess that fails the moment a
    value is "11,160" and a label is "total fee" — so the split is carried in a
    field that already exists rather than re-derived downstream.
    """
    doc = _reader.read(_html())
    metrics = [b for b in doc.body_blocks() if b.role == "metric"]
    assert len(metrics) == 2

    value = "".join(run.text for run in metrics[1].runs if run.bold).strip()
    label = "".join(run.text for run in metrics[1].runs if not run.bold).strip()
    assert value == "11,160"
    assert label == "total fee"


def test_the_warn_tone_survives_the_round_trip():
    doc = _reader.read(_html())
    tones = {b.tone for b in doc.body_blocks() if b.role == "callout"}
    assert tones == {"", "warn"}


# --------------------------------------------------------------------------- #
# Word
# --------------------------------------------------------------------------- #


def _word():
    docx = pytest.importorskip("docx", reason="Word export is an optional extra")
    from artifacts.export.docx import DocxExporter

    return docx.Document(io.BytesIO(DocxExporter().export(_html(), filename="p.docx")))


def test_word_keeps_every_word_of_every_design_block():
    """The loss this suite exists to prevent is silent, so it is checked first."""
    text = "\n".join(p.text for p in _word().paragraphs)
    assert "Reconciliation takes two days a week" in text
    assert "Week 3 is the one that matters" in text
    assert "What it costs" in text
    assert "11,160" in text
    assert "TOTAL FEE" in text
    assert "week 2 becomes a guess" in text


def test_word_sets_a_statement_larger_than_the_body():
    from docx.shared import Pt

    sizes = [
        run.font.size
        for paragraph in _word().paragraphs
        if "Reconciliation takes two days" in paragraph.text
        for run in paragraph.runs
    ]
    assert sizes and all(size == Pt(theme.STATEMENT_PT) for size in sizes)


def test_word_draws_the_rule_and_the_tint_on_a_callout():
    """The bar is the signal and the fill is the nicety.

    Both are asserted because print paths drop shading in some greyscale modes,
    and the rule is what has to survive that.
    """
    from docx.oxml.ns import qn

    for paragraph in _word().paragraphs:
        if "week 2 becomes a guess" not in paragraph.text:
            continue
        properties = paragraph._p.find(qn("w:pPr"))
        assert properties is not None
        borders = properties.find(qn("w:pBdr"))
        assert borders is not None and borders.find(qn("w:left")) is not None
        assert borders.find(qn("w:left")).get(qn("w:color")) == theme.CAUTION.upper()
        shading = properties.find(qn("w:shd"))
        assert shading is not None
        assert shading.get(qn("w:fill")) == theme.WASH_WARN.upper()
        return
    pytest.fail("the warn callout never reached Word")


def test_word_keeps_a_divider_in_the_outline():
    """It stays a Heading 2, so the navigation pane and the PDF bookmark tree
    still see a section. Losing that to make it *look* like a divider would
    trade a real affordance for an appearance.
    """
    headings = [
        p.text.strip() for p in _word().paragraphs if p.style.name.startswith("Heading")
    ]
    assert "What it costs" in headings


# --------------------------------------------------------------------------- #
# PowerPoint
# --------------------------------------------------------------------------- #


def _deck():
    pytest.importorskip("pptx", reason="PowerPoint export is an optional extra")
    from pptx import Presentation

    from artifacts.export.pptx import PptxExporter

    return Presentation(io.BytesIO(PptxExporter().export(_html(), filename="d.pptx")))


def _slide_text(slide) -> str:
    return "\n".join(
        shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
    )


def test_a_statement_gets_a_slide_to_itself():
    """Not a bullet. That is the whole point of the block.

    A statement flattened into a bullet is a statement with its emphasis
    removed, which is the failure the type was added to fix.
    """
    for slide in _deck().slides:
        if "Reconciliation takes two days" in _slide_text(slide):
            assert slide.slide_layout.name == "Blank"
            assert "The problem" not in _slide_text(slide)
            return
    pytest.fail("the statement never reached a slide")


def test_consecutive_metrics_are_one_slide():
    """A row of figures is one slide.

    Four slides each holding a single number is the deck this change exists to
    stop making.
    """
    holding = [s for s in _deck().slides if "TOTAL FEE" in _slide_text(s)]
    assert len(holding) == 1
    assert "18" in _slide_text(holding[0])


def test_a_divider_opens_a_section_without_an_empty_slide_after_it():
    """Found by looking at the output, not at the code.

    The rule that an empty heading is still a section marker was written before
    dividers existed. A divider *is* that marker, so honouring both produced two
    consecutive slides carrying one title with nothing on the second.
    """
    slides = _deck().slides
    titled = [i for i, s in enumerate(slides) if _slide_text(s).strip() == "What it costs"]
    assert len(titled) == 1


def test_a_table_slide_is_named_by_its_section():
    """"Table" is what a renderer says when nobody asked a person.

    A GFM table carries no caption syntax at all, so the section heading is the
    common case rather than the fallback.
    """
    html = render_document(
        title="A proposal",
        blocks=blocks_from_markdown(
            "```divider\nWhat it costs\n```\n\n| Stage | Fee |\n|---|---|\n| Build | 4,960 |\n",
            title="A proposal",
        ),
    )
    from pptx import Presentation

    from artifacts.export.pptx import PptxExporter

    deck = Presentation(io.BytesIO(PptxExporter().export(html, filename="d.pptx")))
    # Positively, because "no slide is called Table" also passes when there is
    # no table slide at all — which would be the worse bug wearing this test's
    # green tick.
    holding = [s for s in deck.slides if any(shape.has_table for shape in s.shapes)]
    assert len(holding) == 1
    assert _slide_text(holding[0]).strip() == "What it costs"


def test_ordinary_prose_is_still_bullets():
    """The change is additive. A document using none of this exports as before."""
    from pptx import Presentation

    from artifacts.export.pptx import PptxExporter

    html = render_document(
        title="Plain", blocks=blocks_from_markdown("## One\n\nA point.\n", title="Plain")
    )
    deck = Presentation(io.BytesIO(PptxExporter().export(html, filename="d.pptx")))
    assert any("A point." in _slide_text(s) for s in deck.slides)


# --------------------------------------------------------------------------- #
# Contrast — the claim `theme.SERIES` makes about itself
# --------------------------------------------------------------------------- #


def _relative_luminance(colour: str) -> float:
    channels = []
    for offset in (0, 2, 4):
        value = int(colour[offset : offset + 2], 16) / 255
        channels.append(
            value / 12.92 if value <= 0.03928 else ((value + 0.055) / 1.055) ** 2.4
        )
    red, green, blue = channels
    return 0.2126 * red + 0.7152 * green + 0.0722 * blue


def _contrast(colour: str, against: str = "ffffff") -> float:
    first, second = _relative_luminance(colour), _relative_luminance(against)
    lighter, darker = max(first, second), min(first, second)
    return (lighter + 0.05) / (darker + 0.05)


@pytest.mark.parametrize("colour", theme.SERIES)
def test_every_series_colour_is_visible_on_paper(colour):
    """WCAG's 3:1 floor for a graphical object.

    Asserted rather than trusted: a contrast claim in a docstring is the kind
    that rots silently, and a palette is the sort of thing somebody adjusts by
    eye on a bright screen.
    """
    assert _contrast(colour) >= 3.0


def test_the_ink_and_the_accent_can_carry_text():
    """4.5:1, which is the floor for body text rather than for a shape."""
    assert _contrast(theme.INK) >= 4.5
    assert _contrast(theme.ACCENT) >= 4.5
    assert _contrast(theme.MUTED) >= 4.5


def test_a_callout_s_text_stays_readable_on_its_own_tint():
    """The panel changes the ground, so the ratio has to be re-checked against
    it rather than against the page.
    """
    assert _contrast(theme.INK, theme.WASH) >= 4.5
    assert _contrast(theme.INK, theme.WASH_WARN) >= 4.5
    assert _contrast(theme.CAUTION, theme.WASH_WARN) >= 3.0


# --------------------------------------------------------------------------- #
# The model is told the vocabulary exists
# --------------------------------------------------------------------------- #


def test_every_brief_offers_the_design_blocks():
    """The leg that decides whether any of the rest is reachable.

    A vocabulary nothing points at is a vocabulary nothing uses, and this file
    would then be a complete, tested, unreachable subsystem — the failure
    `CLAUDE.md` says this codebase has hit fifteen times. `instruction` has
    three branches and one of them was missed on the first pass, which is
    exactly why this is asserted over `BRIEFS` rather than over an example.
    """
    from artifacts.briefs import BRIEFS, instruction

    requests = [f"write a {kind}" for kind in BRIEFS] + ["do something unclassifiable"]
    for request in requests:
        text = instruction(request)
        assert "```statement" in text, request
        assert "```metric" in text, request


def test_the_brief_forbids_inventing_a_figure_to_fill_a_block():
    """Rule 9 in its most dangerous form.

    A block that wants a figure is a standing invitation to produce one, and an
    invented number set large on a slide is the most confident wrong thing this
    product could make.
    """
    from artifacts.briefs import instruction

    text = instruction("write that up as a proposal")
    assert "Never invent a figure" in text
    assert "[to confirm]" in text
