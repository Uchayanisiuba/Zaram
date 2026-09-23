"""A picture in a document survives being exported.

It did not, and nothing reported it. `export/_reader.py` has a small fixed tag
vocabulary — headings, paragraphs, list items, tables — and `img` was not in
it, so every picture in every generated document was dropped on the way to
Word and PowerPoint. The same class of silent loss `_add_table` was written to
end for tables, in the one place nobody had looked, because until 23 September
2026 no document could contain a picture anyway.

That is the trap this file is guarding: a gap that is invisible while the
thing it loses cannot yet be made. `render_chart` embeds a PNG, `create_chart`
exists, and the moment the business layer hands it figures a chart lands in a
document — into an exporter that would have quietly thrown it away.

So the claims here are: the reader sees a picture and where it was; the two
office exporters put it there; a picture that would run off the page is scaled
to fit; the letterhead's logo is chrome and stays out of the prose; and a
remote `src` is never fetched, because an export that makes a network request
is rule 3 broken by a file format.
"""

from __future__ import annotations

import base64
import io

import pytest

from artifacts import export
from artifacts.contracts import Heading, ImageBlock, TableBlock
from artifacts.export import _reader
from artifacts.html import render_document
from artifacts.letterhead import Letterhead


def png(width: int = 400, height: int = 200, colour: str = "#4338ca") -> bytes:
    """A real PNG. python-pptx and python-docx both measure the file itself."""
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), colour).save(buffer, format="PNG")
    return buffer.getvalue()


def data_uri(data: bytes, media_type: str = "image/png") -> str:
    return f"data:{media_type};base64,{base64.b64encode(data).decode('ascii')}"


CHART = png()

DOCUMENT = render_document(
    title="Q3 in one page",
    blocks=[
        Heading(text="What the quarter did", level=2),
        "Revenue held while costs fell.",
        ImageBlock(data=CHART, alt="Revenue by month"),
        Heading(text="What happens next", level=2),
        "Two invoices go out on Monday.",
    ],
)


class TestTheReaderSeesIt:
    def test_a_picture_is_read_out_of_the_markup(self):
        doc = _reader.read(DOCUMENT)

        assert len(doc.images) == 1
        assert doc.images[0].data == CHART
        assert doc.images[0].alt == "Revenue by month"
        assert doc.images[0].media_type == "image/png"

    def test_it_knows_which_section_the_picture_was_in(self):
        doc = _reader.read(DOCUMENT)
        image = doc.images[0]

        # The blocks before it: the h1, the h2, and the paragraph. The picture
        # opened after the third, which is what puts it in the first section
        # and not at the end of the deck.
        before = [b.text.strip() for b in doc.blocks[: image.after_block]]
        assert "Revenue held while costs fell." in before
        assert "What happens next" not in before

    def test_a_remote_source_is_not_fetched(self):
        # Not "is not rendered" — is not *fetched*. An exporter that resolved a
        # URL would put a network request inside a file conversion, where no
        # gate can see it and no log records it.
        doc = _reader.read(
            '<html><body><h1>T</h1>'
            '<img alt="Tracker" src="https://example.com/pixel.png">'
            "</body></html>"
        )

        assert doc.images == []

    def test_a_malformed_picture_does_not_lose_the_prose(self):
        doc = _reader.read(
            '<html><body><h1>T</h1><img src="data:image/png;base64,not base64">'
            "<p>The words are still here.</p></body></html>"
        )

        assert doc.images == []
        assert "The words are still here." in [b.text.strip() for b in doc.blocks]


class TestPowerPoint:
    @staticmethod
    def slides_of(document_html: str):
        from pptx import Presentation

        return Presentation(io.BytesIO(export.render(document_html, "pptx"))).slides

    def test_the_picture_becomes_a_real_picture(self):
        pytest.importorskip("pptx", reason="python-pptx is not installed")
        pictures = [
            shape
            for slide in self.slides_of(DOCUMENT)
            for shape in slide.shapes
            if shape.shape_type == 13  # MSO_SHAPE_TYPE.PICTURE
        ]

        assert pictures, "the picture was dropped on the way to the deck"

    def test_it_lands_in_the_section_it_was_written_in(self):
        pytest.importorskip("pptx", reason="python-pptx is not installed")
        slides = list(self.slides_of(DOCUMENT))
        titles = [s.shapes.title.text if s.shapes.title is not None else "" for s in slides]

        picture_at = next(
            index
            for index, slide in enumerate(slides)
            if any(shape.shape_type == 13 for shape in slide.shapes)
        )

        assert titles.index("What the quarter did") < picture_at
        assert picture_at < titles.index("What happens next")

    def test_the_slide_is_titled_with_the_alt_text(self):
        pytest.importorskip("pptx", reason="python-pptx is not installed")
        titles = [
            slide.shapes.title.text
            for slide in self.slides_of(DOCUMENT)
            if slide.shapes.title is not None
        ]

        assert "Revenue by month" in titles

    def test_a_picture_wider_than_the_slide_is_scaled_to_fit(self):
        pytest.importorskip("pptx", reason="python-pptx is not installed")
        from pptx.util import Inches

        wide = render_document(
            title="Wide",
            blocks=[ImageBlock(data=png(4000, 500), alt="A very wide chart")],
        )
        picture = next(
            shape
            for slide in self.slides_of(wide)
            for shape in slide.shapes
            if shape.shape_type == 13
        )

        assert picture.width <= Inches(13.333) - Inches(1.2)
        assert picture.left >= 0

    def test_the_letterhead_logo_does_not_become_a_slide(self):
        pytest.importorskip("pptx", reason="python-pptx is not installed")
        branded = render_document(
            title="Proposal",
            blocks=["One paragraph."],
            letterhead=Letterhead(name="Zaram", logo=data_uri(png(120, 40))),
        )
        pictures = [
            shape
            for slide in self.slides_of(branded)
            for shape in slide.shapes
            if shape.shape_type == 13
        ]

        assert pictures == [], "the logo is chrome the masthead draws, not a slide"


class TestWord:
    @staticmethod
    def word_of(document_html: str):
        from docx import Document as WordDocument

        return WordDocument(io.BytesIO(export.render(document_html, "docx")))

    def test_the_picture_reaches_the_file(self):
        pytest.importorskip("docx", reason="python-docx is not installed")
        word = self.word_of(DOCUMENT)

        assert word.inline_shapes, "the picture was dropped on the way to Word"

    def test_it_is_never_wider_than_the_page(self):
        pytest.importorskip("docx", reason="python-docx is not installed")
        wide = render_document(
            title="Wide", blocks=[ImageBlock(data=png(4000, 500), alt="Wide")]
        )
        word = self.word_of(wide)
        section = word.sections[0]

        usable = section.page_width - section.left_margin - section.right_margin
        assert word.inline_shapes[0].width <= usable

    def test_the_alt_text_travels_with_it(self):
        pytest.importorskip("docx", reason="python-docx is not installed")
        word = self.word_of(DOCUMENT)

        assert word.inline_shapes[0]._inline.docPr.get("descr") == "Revenue by month"


class TestNothingElseMoved:
    """The tables this change had to walk past on its way in."""

    def test_a_table_still_reaches_word_in_its_place(self):
        pytest.importorskip("docx", reason="python-docx is not installed")
        from docx import Document as WordDocument

        html = render_document(
            title="Fees",
            blocks=[
                Heading(text="Schedule", level=2),
                TableBlock(header=["Item", "Amount"], rows=[["Design", "1,400"]]),
            ],
        )
        word = WordDocument(io.BytesIO(export.render(html, "docx")))

        assert word.tables
        assert word.tables[0].cell(0, 0).text == "Item"
